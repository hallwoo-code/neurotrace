from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"D:\list\A比赛相关\hackthon\NVIDIA\assets\demo-slides\NeuroTrace_演示视频讲解稿.docx")


def set_run_font(run, name="Microsoft YaHei", size=None, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_style_font(style, name="Microsoft YaHei", size=11, bold=False):
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)


def add_labelled_paragraph(doc, label, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    lead = paragraph.add_run(label)
    set_run_font(lead, size=10.5, bold=True)
    body = paragraph.add_run(text)
    set_run_font(body, size=10.5)
    return paragraph


def add_section(doc, timecode, heading, screen, narration):
    paragraph = doc.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(f"{timecode}  {heading}")
    set_run_font(run, size=13, bold=True)
    screen_paragraph = add_labelled_paragraph(doc, "屏幕与操作：", screen)
    screen_paragraph.paragraph_format.keep_with_next = True
    narration_paragraph = add_labelled_paragraph(doc, "讲解：", narration)
    narration_paragraph.paragraph_format.keep_together = True


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    set_style_font(doc.styles["Normal"], size=10.5)
    set_style_font(doc.styles["Title"], size=20, bold=True)
    set_style_font(doc.styles["Heading 1"], size=13, bold=True)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    title_run = title.add_run("NeuroTrace 演示视频讲解稿")
    set_run_font(title_run, size=20, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(11)
    subtitle_run = subtitle.add_run("第三届 DGX Spark 黑客松比赛  Longfuge 团队  建议时长 4 分 30 秒")
    set_run_font(subtitle_run, size=10)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(8)
    intro.paragraph_format.line_spacing = 1.25
    intro_run = intro.add_run(
        "本稿与 NeuroTrace 演示幻灯片对应使用。录制时先展示本次冻结语料运行的判断，再依次操作新建调查、调查工作台、关键核验和案件卷宗。"
        "示例判断必须明确为受语料与证据条件约束的流程结果，不作为固定科研结论。"
    )
    set_run_font(intro_run, size=10.5)

    add_section(
        doc,
        "00:00 至 00:20",
        "开场和团队介绍",
        "展示封面页，停留在项目名称、比赛名称和团队信息上。",
        "大家好，我们是 Longfuge 团队，参加第三届 DGX Spark 黑客松比赛。我是队长吴昊，负责项目统筹、领域证据审核，以及产品和 UI 设计。"
        "梁文豪负责本地服务、数据链路、前端接入、部署与可复跑验证。我们带来的项目是 NeuroTrace。它希望解决的问题很简单：面对一个科研论断，系统不能只给出一个看似确定的答案，而应该保留它的证据条件、反证和可回查路径。",
    )

    add_section(
        doc,
        "00:20 至 00:45",
        "先展示示例结论",
        "切到本次运行结论页，突出“当前证据不足”和“流程演示，非科研结论”。",
        "我们先展示本次运行的受约束判断，再回到过程。推荐演示问题是：情绪唤醒度是否会影响 LPC 振幅？本次在冻结发布语料上的运行显示当前证据不足，并列出适用边界和仍需核验的问题。"
        "这里特别强调，系统根据当前语料与证据条件给出判断；它不是替研究者作出的固定科研结论。NeuroTrace 的核心价值，就是不把有限支持线索包装成已经被证明。",
    )

    add_section(
        doc,
        "00:45 至 01:25",
        "输入问题并启动调查",
        "打开 NeuroTrace 网页的“新建调查”页，输入“情绪唤醒度是否会影响 LPC 振幅？”，确认“运行方式：本地模型 · Ollama”和“资料来源：冻结发布语料”，点击“开始调查”。",
        "接下来进入现场操作。在新建调查页面，我们输入这个研究问题，确认使用本地 Ollama 模型和冻结发布语料。该语料包含 38 份原始 PDF 与 47 张批准证据卡，因此可以展示真实页码、摘录、PDF 哈希和 Evidence Gate；系统同样支持临时上传本地 PDF。"
        "提交后，Trace 从聆听状态进入扫描状态，用户可以看到检索、条件比对、模型分析和 Evidence Gate 等真实进度。若问题缺少会改变结论范围的关键信息，系统只会提出一到两个必要澄清项，而不是把用户带进冗长问卷。",
    )

    # Start the workbench segment on a fresh page. This avoids a visually
    # awkward wrap at the top of page two when the previous section grows.
    doc.add_page_break()
    page_two_lead = doc.add_paragraph()
    page_two_lead.paragraph_format.space_after = Pt(28)
    add_section(
        doc,
        "01:25 至 02:35",
        "讲解三栏调查工作台",
        "进入调查工作台。依次指向左栏的本案基线、中栏的当前判断和证据卡、右栏的真实执行流。只展开两到三张关键证据卡。",
        "调查完成后，我们进入三栏工作台。左侧保留本案基线，也就是最初的问题和补充条件，避免结论在过程中偏离原始研究范围。"
        "中间是当前判断、条件矩阵和证据卡。每张证据卡不仅给出结论，还标注它是支持、限定、反证，还是证据不足，并保留论文信息、页码锚点和关键实验条件。"
        "右侧则是系统的真实执行流，例如资料命中、比对进度和核验状态。我们刻意不展示或伪造模型内部思维链，只呈现实际发生、能够被追踪的系统事件。",
    )

    add_section(
        doc,
        "02:35 至 03:25",
        "展示关键核验和重算",
        "打开关键核验视图，选择一张会影响当前判断的证据卡，说明确认采用、暂缓核验和驳回三种操作。",
        "接下来是关键核验。系统不会要求用户无差别地审阅所有材料，而是按证据角色与核验状态优先呈现可能改变判断的材料：反证、适用边界、条件支持、直接支持和证据不足，并优先显示待核验卡。"
        "对于一张关键证据卡，研究者可以选择确认采用、暂缓核验或驳回。如果选择驳回，必须留下原因；系统会排除该证据并重新计算当前判断，同时保留完整审阅记录。"
        "当前版本不会在驳回后重新生成条件矩阵，因此这里不做该项演示。",
    )

    add_section(
        doc,
        "03:25 至 04:30",
        "卷宗导出和收束",
        "打开案件卷宗，展示导出区域和团队分工页，再回到总览结束演示。",
        "最后，我们打开案件卷宗。卷宗把当前判断、支持证据、边界条件、反证、人工核验记录和导出结果组织在一起。"
        "项目支持 JSON、Markdown 和 CSV 导出：JSON 会保留完整的人工驳回原因，Markdown 与 CSV 当前主要记录审阅状态。导出内容会区分两类信息：一类是从文献中抽取的事实，另一类是基于证据链形成的研究假设或推论建议。"
        "这种分隔能避免把候选判断误写成已经被论文直接证明的结论。总结来说，NeuroTrace 不试图替代研究者，而是帮助研究者把研究直觉变成一条可定位、可核验、可复跑的本地证据链。"
        "它让支持证据、限制条件和反证同时留在结论中。感谢各位评委，我们是 Longfuge 团队。",
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

